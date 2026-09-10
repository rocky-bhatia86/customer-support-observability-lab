from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# Initialize Widescreen Presentation (16:9)
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color Palette Definitions
BG_DARK = RGBColor(9, 13, 22)         # #090d16
CARD_BG = RGBColor(17, 24, 39)        # #111827
CARD_BORDER = RGBColor(31, 41, 55)    # #1f2937
TEXT_WHITE = RGBColor(249, 250, 251)  # #f9fafb
TEXT_MUTED = RGBColor(156, 163, 175) # #9ca3af
CYAN = RGBColor(56, 189, 248)         # #38bdf8
PURPLE = RGBColor(192, 132, 252)      # #c084fc
AMBER = RGBColor(245, 158, 11)        # #f59e0b
ROSE = RGBColor(251, 113, 133)        # #fb7185
GREEN = RGBColor(52, 211, 153)        # #34d399

def set_slide_background(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BG_DARK

def add_header(slide, title_text, subtitle_text, badge_text, badge_color):
    # Title
    txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(9.5), Inches(0.8))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text.upper()
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE
    p.font.name = "Arial"
    
    # Subtitle / Thesis
    p2 = tf.add_paragraph()
    p2.text = f"“{subtitle_text}”"
    p2.font.size = Pt(13)
    p2.font.color.rgb = TEXT_MUTED
    p2.font.name = "Arial"

    # Top-Right Badge
    badge_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.5), Inches(0.45), Inches(2.0), Inches(0.35))
    badge_box.fill.solid()
    badge_box.fill.fore_color.rgb = CARD_BG
    badge_box.line.color.rgb = badge_color
    badge_box.line.width = Pt(1)
    
    p3 = badge_box.text_frame.paragraphs[0]
    p3.text = badge_text.upper()
    p3.font.size = Pt(9)
    p3.font.bold = True
    p3.font.color.rgb = badge_color
    p3.alignment = PP_ALIGN.CENTER

def add_footer(slide, text, badge_text, badge_color):
    # Footer Container
    footer_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.5), Inches(11.733), Inches(0.55))
    footer_box.fill.solid()
    footer_box.fill.fore_color.rgb = CARD_BG
    footer_box.line.color.rgb = CARD_BORDER
    footer_box.line.width = Pt(1)
    
    tf = footer_box.text_frame
    p = tf.paragraphs[0]
    p.text = f"“{text}”"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE

    # Footer Badge
    f_badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.3), Inches(6.6), Inches(2.0), Inches(0.35))
    f_badge.fill.solid()
    f_badge.fill.fore_color.rgb = CARD_BG
    f_badge.line.color.rgb = badge_color
    f_badge.line.width = Pt(1)
    
    p2 = f_badge.text_frame.paragraphs[0]
    p2.text = badge_text.upper()
    p2.font.size = Pt(9)
    p2.font.bold = True
    p2.font.color.rgb = badge_color
    p2.alignment = PP_ALIGN.CENTER


# ==============================================================================
# SLIDE 1: FROM PROTOTYPE TO PRODUCTION
# ==============================================================================
slide_layout = prs.slide_layouts[6] # Blank
slide1 = prs.slides.add_slide(slide_layout)
set_slide_background(slide1)
add_header(slide1, "From Prototype to Production", "What works in a prototype can quietly fail in production.", "Complexity Shift", PURPLE)

# Left: Prototype Card
proto_card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(4.0), Inches(4.7))
proto_card.fill.solid()
proto_card.fill.fore_color.rgb = CARD_BG
proto_card.line.color.rgb = CYAN
proto_card.line.width = Pt(1)

# Proto Title
p_tx = proto_card.text_frame.paragraphs[0]
p_tx.text = "PROTOTYPE"
p_tx.font.size = Pt(12)
p_tx.font.bold = True
p_tx.font.color.rgb = CYAN

# Proto Nodes
proto_nodes = ["INPUT", "SINGLE MODEL", "OUTPUT"]
for i, node_text in enumerate(proto_nodes):
    node = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.3), Inches(2.3 + i*1.0), Inches(3.0), Inches(0.45))
    node.fill.solid()
    node.fill.fore_color.rgb = BG_DARK
    node.line.color.rgb = CYAN
    p = node.text_frame.paragraphs[0]
    p.text = node_text
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_WHITE
    p.alignment = PP_ALIGN.CENTER

# Center Transition
center_tx = slide1.shapes.add_textbox(Inches(5.0), Inches(2.8), Inches(1.5), Inches(2.0))
tf_c = center_tx.text_frame
tf_c.word_wrap = True
p = tf_c.paragraphs[0]
p.text = "SIMPLE\n↓\nCOMPLEXITY\nSHIFT\n↓\nDISTRIBUTED"
p.font.size = Pt(10)
p.font.bold = True
p.font.color.rgb = TEXT_MUTED
p.alignment = PP_ALIGN.CENTER

# Right: Production Card
prod_card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.7), Inches(1.5), Inches(5.833), Inches(4.7))
prod_card.fill.solid()
prod_card.fill.fore_color.rgb = CARD_BG
prod_card.line.color.rgb = PURPLE
prod_card.line.width = Pt(1)

# Prod Nodes
prod_nodes = ["REAL-WORLD TRAFFIC", "ROUTER / ORCHESTRATOR", "SPECIALIST AGENTS (A / B / C)", "TOOLS / APIs (Retry ↺) + DATA", "FINAL RESPONSE"]
for i, p_node_text in enumerate(prod_nodes):
    node = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.2), Inches(2.1 + i*0.8), Inches(4.833), Inches(0.45))
    node.fill.solid()
    node.fill.fore_color.rgb = BG_DARK
    node.line.color.rgb = AMBER if "Retry" in p_node_text else PURPLE
    p = node.text_frame.paragraphs[0]
    p.text = p_node_text
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_WHITE
    p.alignment = PP_ALIGN.CENTER

add_footer(slide1, "Production turns an AI experiment into a system you have to operate.", "OPERATING REALITY", ROSE)


# ==============================================================================
# SLIDE 2: WHAT ACTUALLY GOES WRONG
# ==============================================================================
slide2 = prs.slides.add_slide(slide_layout)
set_slide_background(slide2)
add_header(slide2, "What Actually Goes Wrong", "The biggest AI failures are often hidden inside otherwise successful requests.", "Hidden Failure Modes", ROSE)

# Top Telemetry Bar
bar = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(11.733), Inches(0.55))
bar.fill.solid()
bar.fill.fore_color.rgb = CARD_BG
bar.line.color.rgb = CARD_BORDER
p = bar.text_frame.paragraphs[0]
p.text = "TRACE_ID: tr_9842_x81   |   LATENCY: 4.82s (+3.1s)   |   TOKENS: 14,250 (~3x)   |   STATUS: ✓ 200 OK • SUCCEEDED"
p.font.size = Pt(11)
p.font.color.rgb = GREEN
p.alignment = PP_ALIGN.CENTER

# Trace Flow Box
trace_box = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.2), Inches(11.733), Inches(1.5))
trace_box.fill.solid()
trace_box.fill.fore_color.rgb = CARD_BG
trace_box.line.color.rgb = CARD_BORDER

t_nodes = ["Request", "Agent", "Tool (Stalled)", "Agent Fallback ↩", "Model #2", "Response"]
for i, tn in enumerate(t_nodes):
    n_box = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.1 + i*1.85), Inches(2.7), Inches(1.6), Inches(0.5))
    n_box.fill.solid()
    n_box.fill.fore_color.rgb = BG_DARK
    n_box.line.color.rgb = ROSE if "Stalled" in tn or "Fallback" in tn else CYAN
    p = n_box.text_frame.paragraphs[0]
    p.text = tn
    p.font.size = Pt(10)
    p.font.color.rgb = TEXT_WHITE
    p.alignment = PP_ALIGN.CENTER

# 4 Failure Callouts
callouts = [
    ("TOKEN SPIKE", "Context bloated to 14k tokens due to error loops.", AMBER),
    ("LATENCY", "External API call timed out after 3.8s execution.", AMBER),
    ("RETRY LOOP", "Agent executed 2 extra retries on schema mismatch.", AMBER),
    ("QUALITY REGRESSION", "Semantic relevance score dropped -22% on fallback.", ROSE)
]

for i, (title, desc, color) in enumerate(callouts):
    c_card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i*2.98), Inches(4.0), Inches(2.8), Inches(2.2))
    c_card.fill.solid()
    c_card.fill.fore_color.rgb = CARD_BG
    c_card.line.color.rgb = color
    
    tf = c_card.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = color
    
    p2 = tf.add_paragraph()
    p2.text = desc
    p2.font.size = Pt(10)
    p2.font.color.rgb = TEXT_MUTED

add_footer(slide2, "A successful response does not mean a healthy AI system.", "INTERNAL REALITY", ROSE)


# ==============================================================================
# SLIDE 3: WHY MONITORING ISN'T ENOUGH
# ==============================================================================
slide3 = prs.slides.add_slide(slide_layout)
set_slide_background(slide3)
add_header(slide3, "Why Monitoring Isn’t Enough", "Monitoring tells you something happened. Observability and evaluation tell you what, why, and whether it is getting better.", "System Visibility", PURPLE)

# Col 1: Black Box
col1 = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(3.6), Inches(4.7))
col1.fill.solid()
col1.fill.fore_color.rgb = CARD_BG
col1.line.color.rgb = CARD_BORDER
tf = col1.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "TRADITIONAL MONITORING\n\n\n[ BLACK BOX ]\n\n\nLatency • Errors • Availability\n\n\n“Something failed. But why?”"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = TEXT_MUTED
p.alignment = PP_ALIGN.CENTER

# Col 2: Glass Box
col2 = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.6), Inches(1.5), Inches(5.0), Inches(4.7))
col2.fill.solid()
col2.fill.fore_color.rgb = CARD_BG
col2.line.color.rgb = CYAN

g_nodes = ["AGENT → EXECUTING PLAN", "MODEL → PROMPT & CONTEXT", "TOOL / API → EXECUTION TRACE", "RETRIEVAL → CHUNKS & SCORES"]
for i, gn in enumerate(g_nodes):
    n_box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.9), Inches(2.2 + i*0.9), Inches(4.4), Inches(0.5))
    n_box.fill.solid()
    n_box.fill.fore_color.rgb = BG_DARK
    n_box.line.color.rgb = CYAN
    p = n_box.text_frame.paragraphs[0]
    p.text = gn
    p.font.size = Pt(10)
    p.font.color.rgb = TEXT_WHITE
    p.alignment = PP_ALIGN.CENTER

# Col 3: Evaluation
col3 = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.8), Inches(1.5), Inches(2.733), Inches(4.7))
col3.fill.solid()
col3.fill.fore_color.rgb = CARD_BG
col3.line.color.rgb = PURPLE
tf3 = col3.text_frame
tf3.word_wrap = True
p = tf3.paragraphs[0]
p.text = "EVALUATION LOOP\n\n1. EXECUTION\n2. TRACE\n3. EVALUATION\n4. BASELINE\n5. IMPROVE ↺\n\nOUTPUT QUALITY:\nCorrectness • Relevance"
p.font.size = Pt(11)
p.font.color.rgb = TEXT_MUTED
p.alignment = PP_ALIGN.CENTER

add_footer(slide3, "AI operations require visibility into both execution and intelligence.", "SYSTEM INTELLIGENCE", PURPLE)

# Save the PowerPoint File
output_path = "Prototype_to_Production.pptx"
prs.save(output_path)
print(f"Successfully generated PowerPoint presentation at: {output_path}")
